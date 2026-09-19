from pathlib import Path
import hashlib,json,math,struct
import openpyxl
from pypdf import PdfReader

OUT=Path(__file__).resolve().parent
DATA=OUT.parents[2]/'muju/lab/results/handicap-census-2026-09-14'
C=json.loads((DATA/'census.json').read_text())
WIDTH={0:797,3:3678,4:8012}
BINS={h:(DATA/f'states-h{h}.bin').read_bytes() for h in WIDTH}
R=struct.Struct('<fffBBBBff')
book=openpyxl.load_workbook(OUT/'MHT-handicap-census.xlsx',read_only=True,data_only=True)
assert len(book.sheetnames)==11
errors=[]
for sheet in book:
 for row in sheet.iter_rows():
  for cell in row:
   if cell.data_type=='e':errors.append((sheet.title,cell.coordinate,cell.value))
assert not errors,errors
print('No exported Excel error cells',flush=True)
lookup=book['Position lookup']
assert [lookup[x].value for x in ['B5','B6','B7']]==[4,792,4459]
assert [lookup[x].value for x in ['E5','E6','E7','E8','E9']]==['LEGAL',-3.8,'B clear',9,3]
assert [lookup[x].value for x in ['B11','C11','B15','C15']]==[6,8,0,0]
assert [lookup[x].value for x in ['B20','B23','B25']]==['E6','F5','J10']
assert len(list(book['Black patterns'].iter_rows(min_row=6)))==8012
assert len(list(book['White patterns'].iter_rows(min_row=6)))==797
assert len(list(book['Continuations'].iter_rows(min_row=6)))==83
counts={h:0 for h in WIDTH};illegal={h:0 for h in WIDTH}
for bi,row in enumerate(book['State records'].iter_rows(min_row=6,values_only=True)):
 assert row[0]==bi+1
 for k,h in enumerate([0,3,4]):
  scores,masks=row[1+2*k:3+2*k]
  if bi>=WIDTH[h]:assert scores in [None,''] and masks in [None,''];continue
  assert isinstance(scores,str) and isinstance(masks,str)
  assert scores.startswith("S:") and masks.startswith("M:")
  scores=scores[2:];masks=masks[2:]
  assert len(scores)==3188 and len(masks)==1594
  for wi in range(797):
   v=R.unpack_from(BINS[h],(wi*WIDTH[h]+bi)*24)
   if math.isfinite(v[0]):
    expected=f'{round(v[0]*100)+2000:04d}'
    assert scores[wi*4:wi*4+4]==expected,(h,wi+1,bi+1,'index')
    assert masks[wi*2:wi*2+2]==f'{v[3]:X}{v[4]:X}',(h,wi+1,bi+1,'masks')
    counts[h]+=1
   else:
    assert scores[wi*4:wi*4+4]=='XXXX' and masks[wi*2:wi*2+2]=='XX'
    illegal[h]+=1
assert counts=={0:635203,3:2931337,4:6385505}
assert illegal=={0:6,3:29,4:59}
print('All packed state scores and capture masks match the binary census',counts,flush=True)
book.close()
formulas=openpyxl.load_workbook(OUT/'MHT-handicap-census.xlsx',read_only=True,data_only=False)
assert 'VALUE($B$5)' in formulas['Position lookup']['E12'].value
assert formulas['State records']['B6'].number_format=='@'
formulas.close()
pdf=PdfReader(OUT/'MHT-three-and-four-crystal-openings.pdf')
assert len(pdf.pages)==10
text='\n'.join(p.extract_text() for p in pdf.pages)
for needle in ['2,931,337','6,385,505','19,128','9,316,842','80.03%','83.94%','-3.80','-0.90']:
 assert needle in text,needle
for name,expected in C['baselineHashes'].items():
 assert hashlib.sha256((OUT.parents[2]/'muju'/name).read_bytes()).hexdigest()==expected
out={'exportedFormulaErrors':errors,'packedLegalCounts':counts,'packedCollisionCounts':illegal,'allPackedScoresAndMasksMatch':True,'lookupDefaultVerified':True,'numericAndTextHandicapInputsTestedByBuilder':True,'continuationRows':83,'pdfPages':10,'sourceHashesUnchanged':True,'files':{f.name:{'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in OUT.glob('MHT-*') if f.suffix in ['.pdf','.xlsx','.zip']}}
(OUT/'qa/deliverables-validation.json').write_text(json.dumps(out,indent=2))
print(json.dumps(out,indent=2))
